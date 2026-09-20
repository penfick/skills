"""Layout registry and renderers."""
from __future__ import annotations

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from pptskill.layouts.base import (
    SLIDE_H,
    SLIDE_W,
    add_bg,
    add_footer,
    add_rect,
    add_stat_card,
    add_text,
    add_title_bar,
    blank_slide,
    set_notes,
)
from pptskill.theme import Theme

__all__ = ["render_slide", "LAYOUT_TYPES"]


def _bullets(slide, x, y, w, h, items, theme: Theme, *, size=15, dark=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Inches(0.08)
        p.text = f"•  {item}"
        from pptskill.fonts import style_run

        color = theme.ink if dark else theme.ink_on_light
        style_run(p.runs[0], size=size, color=color, font_cjk=theme.font_cjk, font_latin=theme.font_latin)


def render_cover(prs, cfg, theme, *, page, total):
    s = blank_slide(prs)
    add_bg(s, theme.bg)
    add_rect(s, Inches(0.8), Inches(2.0), Inches(0.9), Inches(0.08), fill=theme.accent)
    add_text(s, Inches(0.8), Inches(2.25), Inches(11), Inches(0.4),
             str(cfg.get("eyebrow", "PRESENTATION")), theme, size=14, bold=True, color=theme.accent, dark=True)
    add_text(s, Inches(0.8), Inches(2.75), Inches(11), Inches(1.2),
             str(cfg["title"]), theme, size=40, bold=True, color=theme.ink, dark=True)
    if cfg.get("subtitle"):
        add_text(s, Inches(0.8), Inches(4.15), Inches(11), Inches(0.6),
                 str(cfg["subtitle"]), theme, size=20, color=theme.muted, dark=True)
    set_notes(s, str(cfg.get("notes", "")))


def render_agenda(prs, cfg, theme, *, page, total):
    s = blank_slide(prs)
    add_bg(s, theme.bg_light)
    add_title_bar(s, str(cfg.get("title", "目录")), theme, subtitle=cfg.get("subtitle"))
    items = cfg.get("items") or []
    for i, item in enumerate(items):
        title = item.get("title", "") if isinstance(item, dict) else str(item)
        desc = item.get("desc", "") if isinstance(item, dict) else ""
        y = Inches(1.7) + Inches(i * 0.95)
        add_rect(s, Inches(0.7), y, Inches(11.9), Inches(0.8),
                 fill=theme.card_on_light, line=theme.border_on_light, radius=True)
        add_text(s, Inches(0.95), y + Inches(0.18), Inches(0.8), Inches(0.5),
                 f"{i+1:02d}", theme, size=22, bold=True, color=theme.accent)
        add_text(s, Inches(2.0), y + Inches(0.2), Inches(3.2), Inches(0.4),
                 title, theme, size=18, bold=True)
        add_text(s, Inches(5.3), y + Inches(0.22), Inches(7), Inches(0.4),
                 desc, theme, size=14, color=theme.muted_on_light)
    add_footer(s, page, total, theme)
    set_notes(s, str(cfg.get("notes", "")))


def render_section(prs, cfg, theme, *, page, total):
    s = blank_slide(prs)
    add_bg(s, theme.bg)
    add_rect(s, Inches(0.8), Inches(2.8), Inches(0.9), Inches(0.08), fill=theme.accent)
    add_text(s, Inches(0.8), Inches(3.05), Inches(11), Inches(1.0),
             str(cfg["title"]), theme, size=36, bold=True, color=theme.ink, dark=True)
    if cfg.get("subtitle"):
        add_text(s, Inches(0.8), Inches(4.2), Inches(11), Inches(0.5),
                 str(cfg["subtitle"]), theme, size=16, color=theme.muted, dark=True)
    set_notes(s, str(cfg.get("notes", "")))


def render_title_body(prs, cfg, theme, *, page, total):
    s = blank_slide(prs)
    add_bg(s, theme.bg_light)
    add_title_bar(s, str(cfg["title"]), theme, subtitle=cfg.get("subtitle"))
    bullets = cfg.get("bullets") or []
    _bullets(s, Inches(0.7), Inches(1.75), Inches(11.5), Inches(4.8), bullets, theme, size=16)
    if cfg.get("callout"):
        add_rect(s, Inches(0.7), Inches(6.1), Inches(11.9), Inches(0.7),
                 fill=theme.accent_soft, radius=True)
        add_text(s, Inches(0.95), Inches(6.25), Inches(11.4), Inches(0.4),
                 str(cfg["callout"]), theme, size=13, bold=True, color=theme.ink_on_light)
    add_footer(s, page, total, theme)
    set_notes(s, str(cfg.get("notes", "")))


def render_two_col(prs, cfg, theme, *, page, total):
    s = blank_slide(prs)
    add_bg(s, theme.bg_light)
    add_title_bar(s, str(cfg.get("title", "")), theme, subtitle=cfg.get("subtitle"))
    for col, key in enumerate(("left", "right")):
        side = cfg.get(key) or {}
        x = Inches(0.7) + Inches(col * 6.15)
        add_rect(s, x, Inches(1.75), Inches(5.9), Inches(4.7),
                 fill=theme.card_on_light, line=theme.border_on_light, radius=True)
        add_text(s, x + Inches(0.3), Inches(1.95), Inches(5.3), Inches(0.4),
                 str(side.get("title", "")), theme, size=18, bold=True)
        _bullets(s, x + Inches(0.3), Inches(2.55), Inches(5.3), Inches(3.5),
                 side.get("items") or [], theme, size=14)
    add_footer(s, page, total, theme)
    set_notes(s, str(cfg.get("notes", "")))


def render_timeline(prs, cfg, theme, *, page, total):
    s = blank_slide(prs)
    add_bg(s, theme.bg_light)
    add_title_bar(s, str(cfg.get("title", "发展脉络")), theme, subtitle=cfg.get("subtitle"))
    milestones = cfg.get("milestones") or []
    add_rect(s, Inches(0.9), Inches(3.15), Inches(11.5), Inches(0.04), fill=theme.border_on_light)
    n = max(len(milestones), 1)
    span = 11.5 / n
    from pptx.enum.shapes import MSO_SHAPE

    for i, m in enumerate(milestones):
        x = Inches(0.85) + Inches(i * span)
        cx = x + Inches(span * 0.45)
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, cx, Inches(3.05), Inches(0.22), Inches(0.22))
        dot.fill.solid()
        dot.fill.fore_color.rgb = theme.accent
        dot.line.fill.background()
        ty = Inches(1.75) if i % 2 == 0 else Inches(3.55)
        year = str(m.get("year", ""))
        title = str(m.get("title", ""))
        desc = str(m.get("desc", ""))
        add_text(s, x, ty, Inches(span), Inches(0.35), year, theme, size=15, bold=True,
                 color=theme.accent, align=PP_ALIGN.CENTER)
        add_text(s, x, ty + Inches(0.38), Inches(span), Inches(0.3), title, theme, size=13,
                 bold=True, align=PP_ALIGN.CENTER)
        add_text(s, x, ty + Inches(0.72), Inches(span), Inches(0.55), desc, theme, size=11,
                 color=theme.muted_on_light, align=PP_ALIGN.CENTER)
    if cfg.get("callout"):
        add_rect(s, Inches(0.7), Inches(5.55), Inches(11.9), Inches(1.0), fill=theme.bg, radius=True)
        add_text(s, Inches(1.0), Inches(5.8), Inches(11.3), Inches(0.55),
                 str(cfg["callout"]), theme, size=15, bold=True, color=theme.ink, dark=True)
    add_footer(s, page, total, theme)
    set_notes(s, str(cfg.get("notes", "")))


def render_data_cards(prs, cfg, theme, *, page, total):
    s = blank_slide(prs)
    add_bg(s, theme.bg_light)
    add_title_bar(s, str(cfg.get("title", "")), theme, subtitle=cfg.get("subtitle"))
    cards = cfg.get("cards") or []
    for i, card in enumerate(cards):
        col, row = i % 2, i // 2
        x = Inches(0.7) + Inches(col * 6.15)
        y = Inches(1.75) + Inches(row * 2.2)
        add_stat_card(s, x, y, Inches(5.9), Inches(1.95),
                      str(card.get("number", "")), str(card.get("label", "")), theme)
    add_footer(s, page, total, theme)
    set_notes(s, str(cfg.get("notes", "")))


def render_closing(prs, cfg, theme, *, page, total):
    s = blank_slide(prs)
    add_bg(s, theme.bg)
    add_rect(s, Inches(0.8), Inches(1.8), Inches(0.9), Inches(0.08), fill=theme.accent)
    add_text(s, Inches(0.8), Inches(2.1), Inches(11), Inches(0.8),
             str(cfg["title"]), theme, size=34, bold=True, color=theme.ink, dark=True)
    lines = cfg.get("lines") or []
    for i, line in enumerate(lines):
        y = Inches(3.3) + Inches(i * 0.55)
        add_rect(s, Inches(0.8), y + Inches(0.12), Inches(0.12), Inches(0.12), fill=theme.accent)
        add_text(s, Inches(1.15), y, Inches(11), Inches(0.4),
                 str(line), theme, size=16, color=theme.ink, dark=True)
    if cfg.get("thanks"):
        add_text(s, Inches(10.2), Inches(6.55), Inches(2.5), Inches(0.3),
                 str(cfg["thanks"]), theme, size=14, color=theme.muted, align=PP_ALIGN.RIGHT, dark=True)
    add_footer(s, page, total, theme, dark=True)
    set_notes(s, str(cfg.get("notes", "")))


def render_chart(prs, cfg, theme, *, page, total, base_dir=None):
    from pptskill.charts import add_chart

    s = blank_slide(prs)
    add_bg(s, theme.bg_light)
    add_title_bar(s, str(cfg["title"]), theme, subtitle=cfg.get("subtitle"))
    add_chart(s, Inches(1.0), Inches(1.75), Inches(11.3), Inches(4.9), cfg.get("chart") or {}, theme)
    add_footer(s, page, total, theme)
    set_notes(s, str(cfg.get("notes", "")))


def render_image_text(prs, cfg, theme, *, page, total, base_dir=None):
    from pptskill.images import add_picture_fit, resolve_image_path

    s = blank_slide(prs)
    add_bg(s, theme.bg_light)
    add_title_bar(s, str(cfg["title"]), theme, subtitle=cfg.get("subtitle"))
    image_cfg = cfg.get("image") or {}
    side = image_cfg.get("side", "left")
    img_path = resolve_image_path(image_cfg.get("path", ""), base_dir)
    # left or right image half
    if side == "left":
        add_picture_fit(s, img_path, Inches(0.7), Inches(1.75), width=Inches(5.8))
        text_x = Inches(6.9)
    else:
        add_picture_fit(s, img_path, Inches(6.8), Inches(1.75), width=Inches(5.8))
        text_x = Inches(0.7)
    _bullets(s, text_x, Inches(1.9), Inches(5.6), Inches(4.5), cfg.get("bullets") or [], theme, size=15)
    add_footer(s, page, total, theme)
    set_notes(s, str(cfg.get("notes", "")))


def render_compare(prs, cfg, theme, *, page, total, base_dir=None):
    s = blank_slide(prs)
    add_bg(s, theme.bg_light)
    add_title_bar(s, str(cfg["title"]), theme, subtitle=cfg.get("subtitle"))
    for col, key in enumerate(("left", "right")):
        side = cfg.get(key) or {}
        highlight = bool(side.get("highlight"))
        x = Inches(0.7) + Inches(col * 6.15)
        fill = theme.accent_soft if highlight else theme.card_on_light
        add_rect(s, x, Inches(1.75), Inches(5.9), Inches(4.7),
                 fill=fill, line=theme.border_on_light, radius=True)
        add_text(s, x + Inches(0.3), Inches(1.95), Inches(5.3), Inches(0.45),
                 str(side.get("title", "")), theme, size=20, bold=True,
                 color=theme.accent if highlight else theme.ink_on_light)
        _bullets(s, x + Inches(0.3), Inches(2.6), Inches(5.3), Inches(3.4),
                 side.get("items") or [], theme, size=14)
    add_footer(s, page, total, theme)
    set_notes(s, str(cfg.get("notes", "")))


def render_kpi(prs, cfg, theme, *, page, total, base_dir=None):
    s = blank_slide(prs)
    add_bg(s, theme.bg_light)
    add_title_bar(s, str(cfg.get("title", "")), theme, subtitle=cfg.get("subtitle"))
    cards = cfg.get("cards") or []
    n = max(len(cards), 1)
    total_w = 11.9
    gap = 0.2
    card_w = (total_w - gap * (n - 1)) / n
    for i, card in enumerate(cards):
        x = Inches(0.7 + i * (card_w + gap))
        add_rect(s, x, Inches(2.0), Inches(card_w), Inches(3.6),
                 fill=theme.card_on_light, line=theme.border_on_light, radius=True)
        add_text(s, x + Inches(0.2), Inches(2.4), Inches(card_w - 0.4), Inches(1.0),
                 str(card.get("number", "")), theme, size=36, bold=True, color=theme.accent,
                 align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.2), Inches(3.6), Inches(card_w - 0.4), Inches(0.5),
                 str(card.get("label", "")), theme, size=16, bold=True,
                 align=PP_ALIGN.CENTER)
        if card.get("hint"):
            add_text(s, x + Inches(0.2), Inches(4.3), Inches(card_w - 0.4), Inches(0.8),
                     str(card["hint"]), theme, size=12, color=theme.muted_on_light,
                     align=PP_ALIGN.CENTER)
    add_footer(s, page, total, theme)
    set_notes(s, str(cfg.get("notes", "")))


def render_quote(prs, cfg, theme, *, page, total, base_dir=None):
    s = blank_slide(prs)
    add_bg(s, theme.bg)
    add_text(s, Inches(1.2), Inches(1.6), Inches(1.5), Inches(1.2),
             "“", theme, size=96, bold=True, color=theme.accent, dark=True)
    add_text(s, Inches(1.4), Inches(2.8), Inches(10.5), Inches(2.0),
             str(cfg.get("text", "")), theme, size=28, bold=True, color=theme.ink, dark=True)
    if cfg.get("cite"):
        add_text(s, Inches(1.4), Inches(5.2), Inches(10.5), Inches(0.5),
                 f"— {cfg['cite']}", theme, size=16, color=theme.muted, dark=True)
    add_footer(s, page, total, theme, dark=True)
    set_notes(s, str(cfg.get("notes", "")))


RENDERERS = {
    "cover": render_cover,
    "agenda": render_agenda,
    "section": render_section,
    "title_body": render_title_body,
    "two_col": render_two_col,
    "timeline": render_timeline,
    "data_cards": render_data_cards,
    "closing": render_closing,
    "chart": render_chart,
    "image_text": render_image_text,
    "compare": render_compare,
    "kpi": render_kpi,
    "quote": render_quote,
}

LAYOUT_TYPES = tuple(sorted(RENDERERS))


def render_slide(prs, cfg: dict, theme: Theme, *, page: int, total: int, base_dir=None) -> None:
    stype = cfg["type"]
    fn = RENDERERS[stype]
    import inspect

    sig = inspect.signature(fn)
    kwargs = {"page": page, "total": total}
    if "base_dir" in sig.parameters:
        kwargs["base_dir"] = base_dir
    fn(prs, cfg, theme, **kwargs)
