"""Structural QA for generated PPTX files."""
from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path

from pptx import Presentation

PLACEHOLDER_RE = re.compile(r"\{\{|\bTODO\b|\bTBD\b|\blorem\b|\bipsum\b", re.IGNORECASE)


@dataclass
class Issue:
    slide: int | None
    level: str  # error | warn
    msg: str
    code: str = ""


@dataclass
class QaReport:
    ok: bool
    issues: list[Issue] = field(default_factory=list)
    pngs: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "ok": self.ok,
            "issues": [asdict(i) for i in self.issues],
            "pngs": list(self.pngs),
        }


def diagnose(
    path: str | Path,
    *,
    visual: bool = False,
    png_dir: str | Path | None = None,
) -> QaReport:
    p = Path(path)
    issues: list[Issue] = []
    pngs: list[str] = []
    if not p.exists():
        return QaReport(ok=False, issues=[Issue(None, "error", f"file not found: {p}")])
    try:
        prs = Presentation(str(p))
    except Exception as e:
        return QaReport(ok=False, issues=[Issue(None, "error", f"cannot open presentation: {e}")])

    if len(prs.slides) == 0:
        issues.append(Issue(None, "error", "presentation has no slides"))

    for idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text or ""
            if PLACEHOLDER_RE.search(text):
                issues.append(Issue(idx, "error", f"placeholder residue in text: {text[:80]!r}", code="placeholder"))
        try:
            notes = slide.notes_slide.notes_text_frame.text or ""
        except Exception:
            notes = ""
        if not notes.strip():
            issues.append(Issue(idx, "warn", "missing speaker notes", code="notes"))

    if visual:
        from pptskill.visual_qa import geometric_issues

        try:
            issues.extend(geometric_issues(p))
        except Exception as e:
            issues.append(Issue(None, "error", f"visual geometric QA failed: {e}", code="visual"))

    if png_dir:
        d = Path(png_dir)
        if d.is_dir():
            pngs = sorted(str(x) for x in d.glob("slide-*.png"))

    ok = not any(i.level == "error" for i in issues)
    return QaReport(ok=ok, issues=issues, pngs=pngs)


def write_report(report: QaReport, path: str | Path) -> None:
    Path(path).write_text(json.dumps(report.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8")
