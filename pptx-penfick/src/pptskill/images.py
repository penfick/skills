"""Local image resolution and embedding."""
from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx.util import Inches


class ImageError(ValueError):
    """Invalid or missing local image."""


def resolve_image_path(path: str | Path, base_dir: str | Path | None) -> Path:
    p = Path(path)
    if not p.is_absolute():
        root = Path(base_dir) if base_dir else Path.cwd()
        p = (root / p).resolve()
    else:
        p = p.resolve()
    if not p.exists():
        raise ImageError(f"image not found: {p}")
    if not p.is_file():
        raise ImageError(f"image is not a file: {p}")
    # validate decodable
    try:
        with Image.open(p) as im:
            im.verify()
    except Exception as e:
        raise ImageError(f"image is not decodable: {p} ({e})") from e
    return p


def add_picture_fit(slide, image_path: Path, x, y, *, width=None, height=None):
    """Embed image preserving aspect ratio. Pass only one of width/height."""
    if width is None and height is None:
        return slide.shapes.add_picture(str(image_path), x, y)
    if width is not None and height is not None:
        # explicit both — only if caller accepts possible distortion; prefer one
        return slide.shapes.add_picture(str(image_path), x, y, width=width, height=height)
    if width is not None:
        return slide.shapes.add_picture(str(image_path), x, y, width=width)
    return slide.shapes.add_picture(str(image_path), x, y, height=height)


def image_size(path: Path) -> tuple[int, int]:
    with Image.open(path) as im:
        return im.size
