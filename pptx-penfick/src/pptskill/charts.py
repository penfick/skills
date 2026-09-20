"""Chart helpers for python-pptx."""
from __future__ import annotations

from typing import Any

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.util import Pt

from pptskill.fonts import style_run
from pptskill.theme import Theme

CHART_KINDS = {"pie", "bar", "line"}


def _chart_type(kind: str) -> XL_CHART_TYPE:
    if kind == "pie":
        return XL_CHART_TYPE.PIE
    if kind == "bar":
        return XL_CHART_TYPE.COLUMN_CLUSTERED
    if kind == "line":
        return XL_CHART_TYPE.LINE_MARKERS
    raise ValueError(f"unsupported chart kind: {kind!r}")


def add_chart(slide, x, y, w, h, chart_cfg: dict[str, Any], theme: Theme):
    kind = str(chart_cfg.get("kind", "pie")).lower()
    if kind not in CHART_KINDS:
        raise ValueError(f"chart.kind must be one of {sorted(CHART_KINDS)}")
    categories = [str(c) for c in (chart_cfg.get("categories") or [])]
    series_list = chart_cfg.get("series") or []
    if not categories:
        raise ValueError("chart.categories must be non-empty")
    if not series_list:
        raise ValueError("chart.series must be non-empty")

    data = CategoryChartData()
    data.categories = categories
    for s in series_list:
        values = list(s.get("values") or [])
        if len(values) != len(categories):
            raise ValueError(
                f"series {s.get('name')!r} values length {len(values)} != categories {len(categories)}"
            )
        data.add_series(str(s.get("name", "series")), values)

    chart = slide.shapes.add_chart(_chart_type(kind), x, y, w, h, data).chart
    chart.has_title = False
    show_legend = bool(chart_cfg.get("show_legend", kind != "pie" or len(series_list) > 1))
    chart.has_legend = show_legend
    if show_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False

    show_values = bool(chart_cfg.get("show_values", True))
    plot = chart.plots[0]
    plot.has_data_labels = show_values
    if show_values:
        labels = plot.data_labels
        labels.show_value = True
        labels.show_category_name = False
        labels.show_series_name = False
        labels.font.size = Pt(11)
        labels.font.bold = True
        if kind == "pie":
            labels.font.color.rgb = theme.ink_on_light
            try:
                labels.position = XL_LABEL_POSITION.OUTSIDE_END
            except Exception:
                pass
        else:
            labels.font.color.rgb = theme.muted_on_light

    # color series / pie slices from theme palette (avoid Office default blue)
    if kind == "pie":
        series = chart.series[0]
        for i in range(len(categories)):
            pt = series.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = theme.series_color(i)
    else:
        for i, series in enumerate(chart.series):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = theme.series_color(i)
            if kind == "line":
                series.format.line.color.rgb = theme.series_color(i)
                series.format.line.width = Pt(2)

    return chart
