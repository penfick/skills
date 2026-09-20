"""Geometric / layout heuristics for visual QA (no renderer required)."""
from __future__ import annotations

from dataclasses import dataclass

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from pptskill.qa import Issue

# slide is 13.333 x 7.5 in = 12192000 x 6858000 EMU
DEFAULT_W = 12192000
DEFAULT_H = 6858000
EMU_PER_IN = 914400


def _shape_issues(slide_idx: int, shape, slide_w: int, slide_h: int) -> list[Issue]:
    issues: list[Issue] = []
    try:
        l, t, w, h = shape.left, shape.top, shape.width, shape.height
    except Exception:
        return issues
    if None in (l, t, w, h):
        return issues

    # fully outside
    if l + w <= 0 or t + h <= 0 or l >= slide_w or t >= slide_h:
        issues.append(Issue(slide_idx, "error", f"shape fully outside slide: {shape.shape_type}"))
        return issues

    tol_x = int(slide_w * 0.02)
    tol_y = int(slide_h * 0.02)
    if l + w > slide_w + tol_x:
        issues.append(Issue(slide_idx, "warn", f"shape overflows right edge: {shape.name}"))
    if t + h > slide_h + tol_y:
        issues.append(Issue(slide_idx, "warn", f"shape overflows bottom edge: {shape.name}"))

    if shape.has_text_frame and (shape.text_frame.text or "").strip():
        if h < int(0.15 * EMU_PER_IN):
            issues.append(Issue(slide_idx, "warn", f"text box very short: {shape.name}"))

    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            img = shape.image
            native = img.size  # (w_px, h_px)
            if native and native[0] and native[1] and w and h:
                native_ar = native[0] / native[1]
                shape_ar = w / h
                if native_ar > 0 and abs(shape_ar - native_ar) / native_ar > 0.15:
                    issues.append(
                        Issue(slide_idx, "warn", f"picture aspect distorted: {shape.name}")
                    )
        except Exception:
            pass
    return issues


def geometric_issues(path) -> list[Issue]:
    prs = Presentation(str(path))
    slide_w = prs.slide_width or DEFAULT_W
    slide_h = prs.slide_height or DEFAULT_H
    issues: list[Issue] = []
    for idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            issues.extend(_shape_issues(idx, shape, slide_w, slide_h))
    return issues
