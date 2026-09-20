from pathlib import Path

from pptx import Presentation

from pptskill.qa import diagnose
from pptskill.render import render_outline
from pptskill.schema import parse_outline


def _outline_dict():
    return {
        "meta": {"title": "Test Deck", "theme": "dark-keynote"},
        "slides": [
            {"type": "cover", "title": "Cover", "subtitle": "Sub", "notes": "n1"},
            {
                "type": "agenda",
                "title": "Agenda",
                "items": [{"title": "A", "desc": "a"}, {"title": "B", "desc": "b"}],
                "notes": "n2",
            },
            {"type": "title_body", "title": "Body", "bullets": ["一", "二"], "notes": "n3"},
            {"type": "section", "title": "Section", "subtitle": "s", "notes": "n3b"},
            {
                "type": "two_col",
                "title": "Cols",
                "left": {"title": "L", "items": ["l1"]},
                "right": {"title": "R", "items": ["r1"]},
                "notes": "n3c",
            },
            {
                "type": "timeline",
                "title": "TL",
                "milestones": [{"year": "2020", "title": "T", "desc": "d"}],
                "notes": "n4",
            },
            {
                "type": "data_cards",
                "title": "Data",
                "cards": [{"number": "1", "label": "one"}],
                "notes": "n5",
            },
            {"type": "closing", "title": "End", "lines": ["bye"], "notes": "n6"},
        ],
    }


def test_render_smoke(tmp_path: Path):
    outline = parse_outline(_outline_dict())
    out = tmp_path / "deck.pptx"
    render_outline(outline, out)
    assert out.exists()
    prs = Presentation(str(out))
    assert len(prs.slides) == 8


def test_qa_ok_on_clean(tmp_path: Path):
    outline = parse_outline(_outline_dict())
    out = tmp_path / "deck.pptx"
    render_outline(outline, out)
    report = diagnose(out)
    assert report.ok
    # notes present → no missing-note errors; warns allowed
    assert not any(i.level == "error" for i in report.issues)


def test_qa_catches_placeholder(tmp_path: Path):
    outline = parse_outline(_outline_dict())
    out = tmp_path / "deck.pptx"
    render_outline(outline, out)
    prs = Presentation(str(out))
    prs.slides[2].shapes[0].text_frame.text = "TODO fill me {{token}}"
    prs.save(str(out))
    report = diagnose(out)
    assert not report.ok


def test_cjk_font_slots(tmp_path: Path):
    outline = parse_outline({
        "meta": {"title": "中文"},
        "slides": [{"type": "title_body", "title": "概况", "bullets": ["中文要点"], "notes": "x"}],
    })
    out = tmp_path / "zh.pptx"
    render_outline(outline, out)
    prs = Presentation(str(out))
    from pptx.oxml.ns import qn

    slide = prs.slides[0]
    found_ea = False
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                rPr = run._r.find(qn("a:rPr"))
                if rPr is None:
                    continue
                ea = rPr.find(qn("a:ea"))
                if ea is not None and ea.get("typeface") == "Microsoft YaHei":
                    found_ea = True
    assert found_ea
