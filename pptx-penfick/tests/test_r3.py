from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from pptskill.qa import diagnose
from pptskill.render import render_outline
from pptskill.schema import parse_outline


def _simple_outline():
    return {
        "meta": {"title": "T"},
        "slides": [
            {"type": "cover", "title": "Hi", "notes": "n"},
            {"type": "title_body", "title": "B", "bullets": ["x"], "notes": "n"},
        ],
    }


def test_visual_ok_on_normal_deck(tmp_path: Path):
    outline = parse_outline(_simple_outline())
    out = tmp_path / "d.pptx"
    render_outline(outline, out)
    report = diagnose(out, visual=True)
    assert report.ok
    # no fully-outside errors
    assert not any(i.level == "error" and "outside" in i.msg for i in report.issues)


def test_visual_flags_overflow(tmp_path: Path):
    outline = parse_outline(_simple_outline())
    out = tmp_path / "d.pptx"
    render_outline(outline, out)
    prs = Presentation(str(out))
    slide = prs.slides[0]
    # push a textbox far past right edge
    tb = slide.shapes.add_textbox(Inches(12), Inches(2), Inches(5), Inches(1))
    tb.text_frame.text = "overflow"
    prs.save(str(out))
    report = diagnose(out, visual=True)
    assert any("overflows right" in i.msg for i in report.issues)


def test_update_slide(tmp_path: Path):
    from pptskill.update import update_slide
    from pptskill.schema import parse_outline

    outline = _simple_outline()
    patched = update_slide(outline, 1, {"title": "New Cover"})
    assert patched["slides"][0]["title"] == "New Cover"
    parse_outline(patched)


def test_update_slide_invalid_index():
    import pytest
    from pptskill.schema import SchemaError
    from pptskill.update import update_slide

    with pytest.raises(SchemaError):
        update_slide(_simple_outline(), 99, {"title": "x"})


def test_assets_check_and_import(tmp_path: Path):
    from PIL import Image
    from pptskill.assets import check_assets, import_asset

    assets = tmp_path / "assets"
    assets.mkdir()
    img = tmp_path / "src.jpg"
    Image.new("RGB", (2000, 1200), (10, 20, 30)).save(img, format="JPEG")
    dest = import_asset(img, assets, name="sample.jpg", max_side=800)
    assert dest.exists()
    from PIL import Image as I
    with I.open(dest) as im:
        assert max(im.size) <= 800

    outline = {
        "meta": {"title": "T"},
        "slides": [{
            "type": "image_text",
            "title": "t",
            "image": {"path": "assets/sample.jpg", "side": "left"},
            "bullets": ["a"],
        }],
    }
    results = check_assets(outline, tmp_path)
    assert results and results[0]["ok"] is True


def test_mcp_handle_line_render(tmp_path: Path):
    import json
    from pptskill.mcp_server import handle_line

    outline_path = tmp_path / "o.json"
    outline_path.write_text(json.dumps(_simple_outline()), encoding="utf-8")
    out = tmp_path / "m.pptx"
    resp = handle_line(json.dumps({
        "id": "1",
        "tool": "render",
        "args": {"outline_path": str(outline_path), "output": str(out)},
    }))
    assert resp["ok"] is True
    assert Path(resp["result"]["path"]).exists()


def test_mcp_unknown_tool():
    from pptskill.mcp_server import handle_line

    resp = handle_line('{"id":"x","tool":"nope","args":{}}')
    assert resp["ok"] is False
