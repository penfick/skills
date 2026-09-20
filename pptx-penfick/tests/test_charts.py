from pathlib import Path

import pytest
from pptx import Presentation

from pptskill.charts import add_chart
from pptskill.layouts.base import blank_slide
from pptskill.render import render_outline
from pptskill.schema import parse_outline
from pptskill.theme import get_theme


def _deck(kind: str, categories, series):
    return {
        "meta": {"title": "T", "theme": "dark-keynote"},
        "slides": [{
            "type": "chart",
            "title": "C",
            "chart": {"kind": kind, "categories": categories, "series": series},
            "notes": "n",
        }],
    }


@pytest.mark.parametrize("kind", ["pie", "bar", "line"])
def test_chart_kinds_render(tmp_path: Path, kind: str):
    if kind == "pie":
        series = [{"name": "s", "values": [1, 2, 3]}]
    else:
        series = [{"name": "s1", "values": [1, 2, 3]}, {"name": "s2", "values": [3, 2, 1]}]
    outline = parse_outline(_deck(kind, ["A", "B", "C"], series))
    out = tmp_path / f"{kind}.pptx"
    render_outline(outline, out)
    prs = Presentation(str(out))
    assert len(prs.slides) == 1


def test_chart_colors_not_office_default():
    theme = get_theme("dark-keynote")
    prs = Presentation()
    prs.slide_width = 12192000
    prs.slide_height = 6858000
    s = blank_slide(prs)
    add_chart(s, 0, 0, 5000000, 4000000, {
        "kind": "pie",
        "categories": ["A", "B"],
        "series": [{"name": "s", "values": [1, 1]}],
    }, theme)
    # theme series[0] is NVIDIA green, not Office blue 4472C4
    c = theme.series_color(0)
    assert (c[0], c[1], c[2]) == (0x76, 0xB9, 0x00)
